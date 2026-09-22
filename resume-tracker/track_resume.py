"""Archive a finished resume PDF and log one row in the Excel tracker.

Columns: Name, Description, Yes/No, Date, Time, Role, Company, PDF, ATS.
Yes/No uses Excel data-validation dropdown (Yes/No) on the column.
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
ARCHIVE_DIR = ROOT / "resume-archive" / "pdf"
TRACKER_DIR = ROOT / "resume-tracker"
WORKBOOK = TRACKER_DIR / "resumes.xlsx"
HEADERS = (
    "Name",
    "Description",
    "Yes/No",
    "Date",
    "Time",
    "Role",
    "Company",
    "PDF",
    "ATS",
)
YES_NO_COL = 3  # 1-based
MAX_ROWS = 5000


def _decision(value: str) -> str:
    key = value.strip().lower()
    if key in {"y", "yes", "true", "1"}:
        return "Yes"
    if key in {"n", "no", "false", "0"}:
        return "No"
    raise SystemExit("Yes/No must be yes or no.")


def _slug(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return (slug or "resume")[:60]


def _one_line(text: str, limit: int = 500) -> str:
    cleaned = " ".join(text.split())
    if len(cleaned) > limit:
        return cleaned[: limit - 3] + "..."
    return cleaned


def _load_openpyxl():
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise SystemExit("Install openpyxl: pip install openpyxl") from None
    return openpyxl, get_column_letter, DataValidation


def _ensure_yes_no_dropdown(ws, DataValidation) -> None:
    """Keep a Yes/No list validation covering the Yes/No column."""
    target = f"C2:C{MAX_ROWS}"
    for dv in list(ws.data_validations.dataValidation):
        ranges = str(dv.sqref) if dv.sqref is not None else ""
        if "C2" in ranges or ranges == target:
            ws.data_validations.dataValidation.remove(dv)
    rule = DataValidation(
        type="list",
        formula1='"Yes,No"',
        allow_blank=False,
        showDropDown=False,  # False = show the dropdown arrow in Excel
        showErrorMessage=True,
    )
    rule.error = "Choose Yes or No from the list"
    rule.errorTitle = "Yes/No"
    rule.prompt = "Select Yes or No"
    rule.promptTitle = "Decision"
    ws.add_data_validation(rule)
    rule.add(target)


def _workbook():
    openpyxl, get_column_letter, DataValidation = _load_openpyxl()
    TRACKER_DIR.mkdir(parents=True, exist_ok=True)
    if WORKBOOK.exists():
        wb = openpyxl.load_workbook(WORKBOOK)
        ws = wb.active
        _ensure_yes_no_dropdown(ws, DataValidation)
        return wb, ws

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumes"
    ws.append(list(HEADERS))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
    widths = (28, 56, 10, 12, 8, 28, 24, 48, 8)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    _ensure_yes_no_dropdown(ws, DataValidation)
    return wb, ws


def _save(wb) -> None:
    tmp = WORKBOOK.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    tmp.replace(WORKBOOK)


def _archive(source: Path, name: str, company: str) -> Path:
    if not source.is_file():
        raise SystemExit(f"PDF not found: {source}")
    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = ARCHIVE_DIR / f"{stamp}-{_slug(name)}-{_slug(company)}.pdf"
    shutil.copy2(source, dest)
    return dest


def cmd_save(args: argparse.Namespace) -> int:
    dest = _archive(Path(args.source), args.name, args.company)
    wb, ws = _workbook()
    now = datetime.now()
    ws.append(
        [
            args.name.strip(),
            _one_line(args.description),
            _decision(args.decision),
            now.strftime("%Y-%m-%d"),
            now.strftime("%H:%M"),
            args.role.strip(),
            args.company.strip(),
            str(dest.relative_to(ROOT)).replace("\\", "/"),
            args.ats if args.ats is not None else "",
        ]
    )
    # Refresh filter range as rows grow
    from openpyxl.utils import get_column_letter

    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(ws.max_row, 1)}"
    _save(wb)
    print(f"archived={dest.relative_to(ROOT)}")
    print(f"tracker={WORKBOOK.relative_to(ROOT)}")
    print(f"row={ws.max_row}")
    print(f"decision={_decision(args.decision)}")
    return 0


def cmd_set(args: argparse.Namespace) -> int:
    if not WORKBOOK.exists():
        raise SystemExit(f"Tracker not found: {WORKBOOK}")
    wb, ws = _workbook()
    if args.row < 2 or args.row > ws.max_row:
        raise SystemExit(f"Row {args.row} is outside 2..{ws.max_row}")
    ws.cell(args.row, YES_NO_COL).value = _decision(args.decision)
    _save(wb)
    print(f"row={args.row}")
    print(f"decision={_decision(args.decision)}")
    return 0


def cmd_list(_: argparse.Namespace) -> int:
    if not WORKBOOK.exists():
        print("rows=0")
        return 0
    _, ws = _workbook()
    print("\t".join(HEADERS))
    count = 0
    for row in ws.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True):
        if not any(row):
            continue
        count += 1
        print("\t".join("" if cell is None else str(cell) for cell in row))
    print(f"rows={count}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Archive a resume PDF and log it.")
    sub = parser.add_subparsers(dest="command", required=True)

    save = sub.add_parser("save", help="Copy PDF into resume-archive and append a row")
    save.add_argument("--name", required=True)
    save.add_argument("--description", required=True)
    save.add_argument("--decision", required=True, help="yes or no")
    save.add_argument("--source", required=True, help="Built PDF path")
    save.add_argument("--role", required=True, help="Target role")
    save.add_argument("--company", required=True, help="Target company")
    save.add_argument("--ats", type=int, default=None)
    save.set_defaults(func=cmd_save)

    set_decision = sub.add_parser("set", help="Change Yes/No on an existing row")
    set_decision.add_argument("--row", type=int, required=True)
    set_decision.add_argument("--decision", required=True)
    set_decision.set_defaults(func=cmd_set)

    listing = sub.add_parser("list", help="Print tracker rows")
    listing.set_defaults(func=cmd_list)

    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
